#!/usr/bin/env python3
"""Build the Variance Research workbook from findings.json + triage.json.

Usage:
  python3 build_research_workbook.py findings.json triage.json \
      --month "03-2026" --out VarianceResearch_032026.xlsx

findings.json is a JSON array of research records (see SKILL.md for the schema).
Items present in triage but absent from findings are carried onto the Research Log
as "below materiality — triage hint only" using their triage hypothesis.
"""
import argparse
import json
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BLUE = "2F5496"
RED = "C00000"
AMBER = "ED7D31"
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")

CAUSE_ORDER = ["TIMING-RECEIPT", "TIMING-SALE", "XFER-UNPOSTED", "UOM-ERROR",
               "SISTER-SKU-MIX", "ADJ-ERROR", "LEDGER-INCONSISTENT", "SHRINK-NORMAL",
               "COUNT-ERROR-SUSPECTED", "DATA-ANOMALY", "UNRESOLVED"]
OPEN_CAUSES = {"COUNT-ERROR-SUSPECTED", "UNRESOLVED", "DATA-ANOMALY"}


def header(ws, row, labels, color=BLUE):
    fill = PatternFill("solid", fgColor=color)
    for c, label in enumerate(labels, 1):
        cell = ws.cell(row=row, column=c, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def autofit(ws, widths):
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def title(ws, text, span):
    ws.cell(row=1, column=1, value=text).font = Font(bold=True, size=14, color=BLUE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("findings")
    ap.add_argument("triage")
    ap.add_argument("--month", default="")
    ap.add_argument("--out", default="VarianceResearch.xlsx")
    args = ap.parse_args()

    with open(args.findings) as fh:
        findings = json.load(fh)
    with open(args.triage) as fh:
        triage = json.load(fh)

    researched_skus = {f.get("sku") for f in findings}
    for it in triage.get("queue", []):
        if it["sku"] in researched_skus:
            continue
        findings.append({
            "sku": it["sku"], "desc": it["desc"], "category": it["cat"],
            "uom": it["uom"], "var_units": it["var_u"],
            "unit_cost": it["unit_cost"], "var_dollars": it["var_dollars"],
            "checks": ["triage only"],
            "evidence": "; ".join(it.get("hints", [])) or "no automated hints",
            "root_cause": (it.get("hypothesis") or "UNRESOLVED").rstrip("?"),
            "confidence": "Low",
            "action": "Below materiality — accept triage hypothesis or batch-review",
            "action_owner": "", "status": "below-materiality",
        })

    findings.sort(key=lambda f: abs(f.get("var_dollars") or 0), reverse=True)
    explained = [f for f in findings
                 if f.get("root_cause") not in OPEN_CAUSES and f.get("status") != "below-materiality"]
    unresolved = [f for f in findings if f.get("root_cause") in OPEN_CAUSES]

    wb = Workbook()

    # ── Dashboard ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Dashboard"
    title(ws, f"Inventory Variance Research — {args.month}", 5)
    total_abs = sum(abs(f.get("var_dollars") or 0) for f in findings)
    explained_abs = sum(abs(f.get("var_dollars") or 0) for f in explained)
    open_abs = sum(abs(f.get("var_dollars") or 0) for f in unresolved)
    rows = [
        ("Items with count variances", len(findings)),
        ("Total absolute variance $", total_abs),
        ("Explained — items / $", f"{len(explained)}  /  ${explained_abs:,.2f}"),
        ("Open (recount or unresolved) — items / $", f"{len(unresolved)}  /  ${open_abs:,.2f}"),
        ("Materiality line", f"${triage.get('materiality', 0):,.0f}"),
    ]
    r = 3
    for label, val in rows:
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=r, column=3, value=val)
        if isinstance(val, float):
            c.number_format = "$#,##0.00"
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="By root cause").font = Font(bold=True, size=12)
    r += 1
    header(ws, r, ["Root cause", "Items", "Variance $ (net)", "Variance $ (abs)"])
    by_cause = defaultdict(list)
    for f in findings:
        by_cause[f.get("root_cause") or "UNRESOLVED"].append(f)
    for cause in CAUSE_ORDER + sorted(set(by_cause) - set(CAUSE_ORDER)):
        if cause not in by_cause:
            continue
        grp = by_cause[cause]
        r += 1
        ws.cell(row=r, column=1, value=cause)
        ws.cell(row=r, column=2, value=len(grp))
        net = ws.cell(row=r, column=3, value=sum(f.get("var_dollars") or 0 for f in grp))
        ab = ws.cell(row=r, column=4, value=sum(abs(f.get("var_dollars") or 0) for f in grp))
        net.number_format = ab.number_format = "$#,##0.00"
    autofit(ws, [42, 10, 18, 18])
    ws.sheet_properties.tabColor = BLUE

    # ── Research Log ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Research Log")
    cols = ["SKU", "Description", "Category", "UOM", "Var Units", "Unit Cost",
            "Var $", "Root Cause", "Confidence", "Evidence", "Checks Performed",
            "Recommended Action", "Owner", "Status"]
    header(ws, 1, cols)
    for r, f in enumerate(findings, 2):
        vals = [f.get("sku"), f.get("desc"), f.get("category"), f.get("uom"),
                f.get("var_units"), f.get("unit_cost"), f.get("var_dollars"),
                f.get("root_cause"), f.get("confidence"), f.get("evidence"),
                "; ".join(f.get("checks", [])), f.get("action"),
                f.get("action_owner"), f.get("status")]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 5:
                cell.number_format = "#,##0.0"
            elif c == 6:
                cell.number_format = "$#,##0.0000"
            elif c == 7:
                cell.number_format = "$#,##0.00"
        cause = f.get("root_cause")
        fill = (RED_FILL if cause in ("UNRESOLVED", "COUNT-ERROR-SUSPECTED")
                else AMBER_FILL if f.get("status") == "below-materiality"
                else GREEN_FILL if f.get("confidence") == "High" else None)
        if fill:
            ws.cell(row=r, column=8).fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(findings) + 1}"
    autofit(ws, [11, 34, 12, 6, 11, 11, 12, 20, 11, 48, 34, 40, 16, 16])
    ws.sheet_properties.tabColor = "70AD47"

    # ── Action Items ───────────────────────────────────────────────────────
    ws = wb.create_sheet("Action Items")
    title(ws, "Corrective Actions — forward to branches", 6)
    r = 3
    by_action_cause = defaultdict(list)
    for f in explained:
        by_action_cause[f.get("root_cause")].append(f)
    for cause in CAUSE_ORDER:
        grp = by_action_cause.get(cause)
        if not grp:
            continue
        ws.cell(row=r, column=1, value=cause).font = Font(bold=True, size=12, color=BLUE)
        r += 1
        header(ws, r, ["SKU", "Description", "Var Units", "Var $", "Action", "Owner"], AMBER)
        for f in grp:
            r += 1
            for c, v in enumerate([f.get("sku"), f.get("desc"), f.get("var_units"),
                                   f.get("var_dollars"), f.get("action"),
                                   f.get("action_owner")], 1):
                cell = ws.cell(row=r, column=c, value=v)
                if c == 3:
                    cell.number_format = "#,##0.0"
                elif c == 4:
                    cell.number_format = "$#,##0.00"
        r += 2
    autofit(ws, [11, 34, 11, 12, 52, 18])
    ws.sheet_properties.tabColor = AMBER

    # ── Unresolved & Recounts ──────────────────────────────────────────────
    ws = wb.create_sheet("Unresolved & Recounts")
    cols = ["SKU", "Description", "Category", "Var Units", "Var $", "Disposition",
            "What to verify", "Checks already performed"]
    header(ws, 1, cols, RED)
    for r, f in enumerate(unresolved, 2):
        for c, v in enumerate([f.get("sku"), f.get("desc"), f.get("category"),
                               f.get("var_units"), f.get("var_dollars"),
                               f.get("root_cause"), f.get("action"),
                               "; ".join(f.get("checks", []))], 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c == 4:
                cell.number_format = "#,##0.0"
            elif c == 5:
                cell.number_format = "$#,##0.00"
    ws.freeze_panes = "A2"
    autofit(ws, [11, 34, 12, 11, 12, 24, 44, 40])
    ws.sheet_properties.tabColor = RED

    wb.save(args.out)
    print(f"Saved {args.out}")
    print(f"  Researched/logged: {len(findings)} items | explained: {len(explained)} "
          f"(${explained_abs:,.2f}) | open: {len(unresolved)} (${open_abs:,.2f})")


if __name__ == "__main__":
    main()
